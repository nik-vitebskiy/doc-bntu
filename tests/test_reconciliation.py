from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import AuditLog, OrderItem, Organization
from app.services.organization_service import save_item
from app.services.reconciliation_service import reconcile_xlsx, report_xlsx


def write_ais_file(path, organization_name):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Факультет", "Организация-заказчик", "Код специальности, направления специальности, специализации", 2027,
    ])
    sheet.append(["Тестовый факультет", organization_name, "S-1", 15])
    sheet.append(["Тестовый факультет", organization_name, "S-2", 4])
    sheet.append(["Тестовый факультет", "Чужая организация", "S-9", 99])
    workbook.save(path)


def test_reconciliation_reports_three_difference_types_without_mutation(session, user, actor, organization, contract, tmp_path):
    faculty_id = contract.faculty_links[0].faculty_id
    save_item(session, contract.id, "S-1", "", {"faculty_id": faculty_id, "demand_2027": "10"}, audit_actor=actor)
    save_item(session, contract.id, "S-3", "", {"faculty_id": faculty_id, "demand_2027": "3"}, audit_actor=actor)
    session.flush()
    before = [(item.specialty, item.demand_json) for item in contract.items]
    path = tmp_path / "ais.xlsx"
    write_ais_file(path, organization.short_name)

    report = reconcile_xlsx(session, path, "АИС.xlsx", audit_actor=actor)

    assert report.organizations_checked == 1
    assert report.organizations_skipped == 1
    assert report.skipped_names == ["Чужая организация"]
    assert {(row.kind, row.specialty, row.ours, row.ais) for row in report.differences} == {
        ("changed", "S-1", 10, 15),
        ("ais_only", "S-2", 0, 4),
        ("system_only", "S-3", 3, 0),
    }
    session.expire_all()
    assert [(item.specialty, item.demand_json) for item in contract.items] == before
    assert session.scalar(select(func.count()).select_from(Organization)) == 1
    audit = session.scalar(select(AuditLog).where(AuditLog.entity_type == "ais_reconciliation"))
    assert audit and "3 расхождений, 1 организаций проверено" in audit.entity_label

    stream = report_xlsx(report)
    rows = list(load_workbook(BytesIO(stream.read()), data_only=True).active.values)
    assert len(rows) == 4
    assert {row[0] for row in rows[1:]} == {"Значения отличаются", "Есть в АИС, нет у нас", "Есть у нас, нет в АИС"}


def test_reconciliation_web_report_and_export(client, session, actor, organization, contract, tmp_path):
    faculty_id = contract.faculty_links[0].faculty_id
    save_item(session, contract.id, "S-1", "", {"faculty_id": faculty_id, "demand_2027": "10"}, audit_actor=actor)
    path = tmp_path / "ais.xlsx"
    write_ais_file(path, organization.short_name)
    with path.open("rb") as source:
        response = client.post(
            "/reconciliation", files={"file": ("АИС.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 200
    assert "Значения отличаются" in response.text
    export_url = response.text.split('/reconciliation/')[1].split('/export')[0]
    export = client.get(f"/reconciliation/{export_url}/export")
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats")

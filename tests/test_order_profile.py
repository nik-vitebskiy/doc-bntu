from datetime import date

from openpyxl import Workbook
from sqlalchemy import select

from app.models import Contract
from app.services.application_service import create_application, save_application_item
from app.services.import_service import import_xlsx
from app.services.organization_service import (
    compare_agreement_order,
    register_additional_agreement,
    save_item,
)


def test_profile_is_saved_manually_and_rendered_in_order_grid(
    session, contract, user, actor, client
):
    item = save_item(
        session,
        contract.id,
        "PROFILE-01",
        "Инженер",
        {"profile": "Проектирование машин", "demand_2027": "3"},
        user_id=user.id,
        audit_actor=actor,
    )

    assert item.profile == "Проектирование машин"
    page = client.get(f"/organizations/{contract.organization_id}")
    assert page.status_code == 200
    assert "Профилизация" in page.text
    assert 'name="profile" value="Проектирование машин"' in page.text
    assert 'class="order-table"' in page.text


def test_application_profile_is_saved_and_rendered(
    session, organization, user, actor, client
):
    application = create_application(
        session,
        organization.id,
        ["Автотракторный"],
        "PROFILE-APP",
        "",
        "",
        user.id,
        audit_actor=actor,
    )
    item = save_application_item(
        session,
        application,
        "PROFILE-APP-01",
        "Инженер",
        {"profile": "Автоматизированные системы", "demand_2028": "2"},
        audit_actor=actor,
    )

    assert item.profile == "Автоматизированные системы"
    page = client.get(f"/applications/{application.id}")
    assert page.status_code == 200
    assert 'name="profile" value="Автоматизированные системы"' in page.text
    assert 'class="order-table"' in page.text


def test_order_item_position_is_stable_after_api_update(session, contract, user, actor, client):
    first = save_item(
        session, contract.id, "STABLE-01", "Инженер",
        {"profile": "Первый", "demand_2027": "1"},
        user_id=user.id, audit_actor=actor,
    )
    second = save_item(
        session, contract.id, "STABLE-02", "Инженер",
        {"profile": "Второй", "demand_2027": "2"},
        user_id=user.id, audit_actor=actor,
    )
    before = client.get(f"/organizations/{contract.organization_id}").text
    before_order = [before.index(f'id="item-{item.id}"') for item in (first, second)]

    response = client.put(f"/api/order-items/{first.id}", json={
        "faculty_id": first.faculty_id,
        "specialty_id": first.specialty_id,
        "profile": "Первый изменён",
        "qualification": "Инженер",
        "years": {"2027": 3},
    })
    assert response.status_code == 200

    after = client.get(f"/organizations/{contract.organization_id}").text
    after_order = [after.index(f'id="item-{item.id}"') for item in (first, second)]
    assert before_order[0] < before_order[1]
    assert after_order[0] < after_order[1]
    assert "из справочника:" not in after


def test_profile_is_copied_and_visible_when_revision_changes(
    session, contract, user, actor
):
    source = save_item(
        session,
        contract.id,
        "PROFILE-02",
        "Инженер",
        {"profile": "Исходный профиль", "demand_2027": "1"},
        user_id=user.id,
        audit_actor=actor,
    )
    agreement = register_additional_agreement(
        session, contract, "1", date(2026, 9, 29), user.id, audit_actor=actor
    )
    copied = agreement.orders[0].items[0]
    assert copied.profile == source.profile

    save_item(
        session,
        contract.id,
        copied.specialty,
        copied.qualification,
        {
            "faculty_id": str(copied.faculty_id),
            "profile": "Изменённый профиль",
            "demand_2027": "1",
        },
        item=copied,
        user_id=user.id,
        audit_actor=actor,
    )
    rows, _years = compare_agreement_order(session, agreement)
    changed = next(row for row in rows if row["specialty"] == "PROFILE-02")
    assert changed["change"] == "Изменено"
    assert changed["before"]["profile"] == "Исходный профиль"
    assert changed["after"]["profile"] == "Изменённый профиль"


def test_import_without_profile_column_preserves_manual_profile(session, user, tmp_path):
    path = tmp_path / "without-profile.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Организация-заказчик", "УНП", "Факультет", "Номер договора",
        "Статус", "Код специальности, направления специальности, специализации", "2027",
    ])
    sheet.append([
        "Профильный завод", "100000099", "Автотракторный",
        "№P-1 от 01.01.2026", "Активен", "PROFILE-03", 2,
    ])
    workbook.save(path)

    import_xlsx(session, path, user.id, create_unknown_organizations=True)
    contract = session.scalar(select(Contract).where(Contract.number == "P-1"))
    contract.items[0].profile = "Ручное значение"
    session.flush()

    import_xlsx(session, path, user.id, create_unknown_organizations=True)
    session.expire_all()
    assert contract.items[0].profile == "Ручное значение"

    sheet.insert_cols(7)
    sheet.cell(1, 7).value = "Профилизация"
    sheet.cell(2, 7).value = "Значение из файла"
    workbook.save(path)
    import_xlsx(session, path, user.id, create_unknown_organizations=True)
    session.expire_all()
    assert contract.items[0].profile == "Значение из файла"

from datetime import date, datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.models import AuditLog
from app.services.application_service import create_application, save_application_item
from app.services.organization_service import create_contract, create_organization, save_item


def _rows(response):
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    return list(workbook.active.iter_rows(values_only=True))


def test_contract_export_uses_registry_filters_and_full_export_has_all(
    session, organization, contract, user, actor, client
):
    other = create_organization(session, name="Вторая организация", unp="999000002", audit_actor=actor)
    closed = create_contract(
        session, other.id, ["Другой факультет"], "CLOSED-02", "2032-12-31",
        start_date="2026-01-01", audit_actor=actor,
    )
    closed.status = "Закрыт"
    session.commit()
    save_item(session, contract.id, "EXP-01", "", {
        "faculty_id": str(contract.faculty_links[0].faculty_id), "demand_2027": "1",
    }, user_id=user.id, audit_actor=actor)

    full = client.get("/export/contracts")
    filtered = client.get("/export/contracts", params={
        "status": "Активен", "q": contract.number,
        "faculty": contract.faculty_links[0].faculty.name,
        "end_year": str(contract.end_date.year), "urgency": "later",
    })
    assert full.status_code == filtered.status_code == 200
    assert len(_rows(full)) == 3
    filtered_rows = _rows(filtered)
    assert len(filtered_rows) == 2
    assert contract.number in filtered_rows[1][1]
    assert filtered_rows[1][6] == "EXP-01"


def test_application_export_uses_filters_and_contains_specialty_codes(
    session, organization, user, actor, client
):
    active_end = date.today() + timedelta(days=10)
    closed_end = date.today() + timedelta(days=100)
    active = create_application(session, organization.id, ["Факультет заявок"], "APP-A", "2026-09-01", active_end.isoformat(), user.id, audit_actor=actor)
    closed = create_application(session, organization.id, ["Факультет заявок"], "APP-C", "2026-09-02", closed_end.isoformat(), user.id, audit_actor=actor)
    closed.status = "Закрыт"
    session.commit()
    save_application_item(session, active, "APP-SPEC", "", {
        "faculty_id": str(active.faculty_links[0].faculty_id), "demand_2027": "2",
    }, audit_actor=actor)

    full = _rows(client.get("/export/applications"))
    filtered = _rows(client.get("/export/applications", params={
        "status": "Заявка", "faculty": "Факультет заявок", "urgency": "due_30",
    }))
    assert len(full) == 3
    assert len(filtered) == 2
    assert filtered[1][0] == "APP-A"
    assert filtered[1][7] == "APP-SPEC"


def test_import_export_page_has_history_controls_and_registry_is_clean(client, session, user):
    session.add(AuditLog(
        user_id=user.id, action="FILE_UPLOAD", entity_type="excel_import",
        entity_label="Импорт Excel history.xlsx",
        diff={"old": {}, "new": {"filename": "history.xlsx", "rows_processed": 12, "organizations": 2, "contracts_created": 1}},
        timestamp=datetime(2026, 9, 30, 9, 15, tzinfo=timezone.utc),
    ))
    session.add(AuditLog(
        user_id=user.id, action="FILE_UPLOAD", entity_type="excel_import",
        entity_label="Импорт Excel unchanged.xlsx",
        diff={"old": {}, "new": {
            "filename": "unchanged.xlsx", "rows_processed": 0,
            "organizations_created": 0, "organizations_updated": 0,
            "contracts_created": 0, "contracts_updated": 0,
            "applications_created": 0, "applications_updated": 0,
            "order_items_created": 0, "order_items_updated": 0,
        }},
        timestamp=datetime(2026, 9, 30, 9, 16, tzinfo=timezone.utc),
    ))
    session.commit()

    page = client.get("/import-export")
    assert page.status_code == 200
    assert "history.xlsx" in page.text
    assert "Строк: 12" in page.text
    assert "договоры — создано: 1, обновлено: 0" in page.text
    assert "unchanged.xlsx" in page.text and "Изменений нет." in page.text
    assert 'action="/import"' in page.text and 'action="/reconciliation"' in page.text
    assert 'href="/export/contracts"' in page.text and 'href="/export/applications"' in page.text
    assert 'data-persist-checkbox="import-export:create-organizations:user:' in page.text

    registry = client.get("/")
    assert 'action="/import"' not in registry.text
    assert 'action="/reconciliation"' not in registry.text


def test_registry_export_links_preserve_active_filters(client):
    contracts = client.get("/contracts", params={"q": "ABC", "status": "Активен", "end_year": "2030", "urgency": "due_30"})
    applications = client.get("/applications", params={"q": "APP", "status": "Заявка", "urgency": "due_30"})
    assert "/export/contracts?" in contracts.text
    assert "q=ABC" in contracts.text and "end_year=2030" in contracts.text and "urgency=due_30" in contracts.text
    assert "/export/applications?" in applications.text
    assert "q=APP" in applications.text and "urgency=due_30" in applications.text

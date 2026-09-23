from pathlib import Path

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.models import AppSetting, AppUser, AuditLog
from app.services.admin_service import get_bntu_requisites, update_bntu_requisites
from app.services.auth_service import hash_password
from app.services.document_service import render_agreement


def _document_text(path: Path) -> str:
    document = WordDocument(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def _settings_form(**overrides):
    values = {
        "full_name": "Белорусский национальный технический университет",
        "signer_position": "Проректор по учебной работе",
        "signer_name": "Николайчик Юрий Александрович",
        "power_of_attorney_number": "01-19/1090",
        "power_of_attorney_date": "2026-02-10",
        "legal_address": "220013, г. Минск, пр-т Независимости, 65",
        "unp": "100 354 447",
        "okpo": "02 071 903",
        "bank_account": "BY69 AKBB 3632 9016 3601 3550 0000",
        "bank_name": "ОАО «АСБ Беларусбанк»",
        "bic": "AKBBBY2X",
    }
    values.update(overrides)
    return values


def test_admin_updates_requisites_and_change_is_audited(client, session, actor):
    update_bntu_requisites(session, _settings_form(), audit_actor=actor)

    response = client.post("/settings", data=_settings_form(power_of_attorney_number="NEW-2026/77"))

    assert response.status_code == 303
    session.expire_all()
    setting = session.get(AppSetting, "bntu.power_of_attorney_number")
    assert setting.value == "NEW-2026/77"
    event = session.scalars(select(AuditLog).where(
        AuditLog.entity_type == "app_setting",
        AuditLog.action == "UPDATE",
        AuditLog.entity_label == "Настройка: Номер доверенности",
    )).one()
    assert event.diff == {
        "old": {"value": "01-19/1090"},
        "new": {"value": "NEW-2026/77"},
    }


def test_head_cannot_open_settings_and_has_no_menu_link(session):
    from app.main import app

    head = AppUser(
        username="settings-head",
        password_hash=hash_password("Head-password-123"),
        full_name="Руководитель",
        role="HEAD",
        must_change_password=False,
    )
    session.add(head)
    session.commit()
    with TestClient(app, follow_redirects=False) as browser:
        assert browser.post("/login", data={
            "username": "settings-head", "password": "Head-password-123",
        }).status_code == 303
        assert browser.get("/settings").status_code == 403
        assert 'href="/settings"' not in browser.get("/").text


def test_agreement_uses_current_requisites_and_preserves_existing_file(session, contract, actor, tmp_path):
    original = _settings_form(power_of_attorney_number="OLD-POWER")
    update_bntu_requisites(session, original, audit_actor=actor)
    old_path = tmp_path / "old.docx"
    render_agreement(contract, old_path, get_bntu_requisites(session))
    old_bytes = old_path.read_bytes()
    assert "доверенности №OLD-POWER от 10.02.2026" in _document_text(old_path)

    updated = _settings_form(power_of_attorney_number="NEW-POWER")
    update_bntu_requisites(session, updated, audit_actor=actor)

    assert old_path.read_bytes() == old_bytes
    assert "доверенности №OLD-POWER от 10.02.2026" in _document_text(old_path)
    new_path = tmp_path / "new.docx"
    render_agreement(contract, new_path, get_bntu_requisites(session))
    assert "доверенности №NEW-POWER от 10.02.2026" in _document_text(new_path)


def test_agreement_with_empty_requisites_uses_placeholders(session, contract, tmp_path):
    assert get_bntu_requisites(session) == {
        "full_name": "", "signer_position": "", "signer_name": "",
        "power_of_attorney_number": "", "power_of_attorney_date": "",
        "legal_address": "", "unp": "", "okpo": "", "bank_account": "",
        "bank_name": "", "bic": "",
    }
    output = tmp_path / "empty.docx"
    render_agreement(contract, output, get_bntu_requisites(session))
    text = _document_text(output)
    assert "доверенности №___ от ___" in text
    assert "УНП ___, ОКПО ___" in text
    assert "р/с ___" in text

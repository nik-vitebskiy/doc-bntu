import json
from datetime import datetime, date
from openpyxl import load_workbook
from .models import Organization, Contract, OrderItem

def text(v): return "" if v is None else str(v).strip()
def parse_date(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try: return datetime.strptime(text(value), fmt).date()
        except ValueError: pass
    return None

def import_xlsx(path, db):
    ws = load_workbook(path, data_only=True).active
    headers = [text(c.value) for c in ws[1]]
    ix = {h: i for i, h in enumerate(headers)}
    def field(row, label): return row[ix[label]] if label in ix else None
    years = [(int(h), i) for i, h in enumerate(headers) if text(h).isdigit() and 2000 <= int(h) <= 2100]
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = text(field(row, "Организация-заказчик"))
        if not name: continue
        org = db.query(Organization).filter_by(name=name).first()
        if not org:
            org = Organization(name=name, full_name=text(field(row, "Полное наименование")), address=text(field(row, "Адрес юридический")), profile=text(field(row, "Профиль образования")), phones=text(field(row, "Телефоны")), department=text(field(row, "Ведомство")), contact=text(field(row, "Директор")))
            db.add(org); db.flush()
        faculty, number = text(field(row, "Факультет")), text(field(row, "Номер договора"))
        contract = db.query(Contract).filter_by(organization_id=org.id, faculty=faculty, number=number).first()
        if not contract:
            contract = Contract(organization_id=org.id, faculty=faculty, number=number, status=text(field(row, "Статус")) or "Активен", start_date=parse_date(field(row, "Дата начала договора")), end_date=parse_date(field(row, "Дата окончания договора")))
            db.add(contract); db.flush()
        demand = {str(year): (row[i] or 0) for year, i in years}
        specialty = text(field(row, "Код специальности, направления специальности, специализации"))
        qualification = text(field(row, "Квалификация"))
        if specialty and not db.query(OrderItem).filter_by(contract_id=contract.id, specialty=specialty, qualification=qualification).first():
            db.add(OrderItem(contract_id=contract.id, specialty=specialty, qualification=qualification, demand_json=json.dumps(demand, ensure_ascii=False)))
        count += 1
    db.commit()
    return count

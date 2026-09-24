"""Importer for manually uploaded Excel exports."""

from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func

from ..models import AnnualDemand, Contract, ContractFaculty, OrderItem, Organization
from .audit_service import AuditAction, audited, current_audit_batch
from .organization_service import get_or_create_faculty, get_or_create_order, get_or_create_specialty


@dataclass(frozen=True)
class ImportResult:
    rows_processed: int
    organizations: int
    contracts: int
    faculties: int
    specialties: int
    organizations_created: int
    contracts_created: int
    faculty_links_created: int
    order_items_created: int

    def log_line(self, source: str) -> str:
        return (
            f"Импорт {source}: загружено {self.contracts} договоров, "
            f"{self.organizations} организаций, {self.faculties} факультетов, "
            f"{self.specialties} специальностей из {self.rows_processed} строк; "
            f"новых: {self.contracts_created} договоров, {self.organizations_created} организаций, "
            f"{self.faculty_links_created} связей с факультетами, {self.order_items_created} строк заказа"
        )


def text(value):
    return "" if value is None else str(value).strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text(value), fmt).date()
        except ValueError:
            pass
    return None


def _find_organization(session, unp, name):
    if unp:
        return session.query(Organization).filter_by(unp=unp).first(), unp
    existing = session.query(Organization).filter(func.lower(Organization.short_name) == name.lower()).first()
    if existing:
        return existing, existing.unp
    generated = f"9{int(sha256(name.casefold().encode()).hexdigest(), 16) % 100_000_000:08d}"
    if session.query(Organization.id).filter_by(unp=generated).first():
        raise ValueError(f"Неоднозначный технический УНП для «{name}».")
    return None, generated


@audited
def import_xlsx(session, path: Path, user_id=None, original_filename=None) -> ImportResult:
    """Upsert source details; keep statuses set manually by staff."""
    worksheet = load_workbook(path, data_only=True, read_only=True).active
    headers = [text(cell.value) for cell in worksheet[1]]
    index = {header: i for i, header in enumerate(headers)}
    required = {"Организация-заказчик", "Факультет", "Номер договора", "Код специальности, направления специальности, специализации"}
    missing = sorted(required - index.keys())
    if missing:
        raise ValueError(f"В Excel отсутствуют столбцы: {', '.join(missing)}")

    def field(row, label):
        return row[index[label]] if label in index else None

    years = [(int(header), i) for i, header in enumerate(headers) if header.isdigit() and 2000 <= int(header) <= 2100]
    seen_organizations, seen_contracts, seen_faculties, seen_specialties = set(), set(), set(), set()
    rows_processed = organizations_created = contracts_created = faculty_links_created = order_items_created = 0

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        name = text(field(row, "Организация-заказчик"))
        if not name:
            continue
        organization, unp = _find_organization(session, text(field(row, "УНП")), name)
        full_name = text(field(row, "Полное наименование")) or name
        address = text(field(row, "Адрес юридический")) or None
        authority = text(field(row, "Ведомство")) or None
        phone = text(field(row, "Телефоны")) or None
        if organization is None:
            organization = Organization(unp=unp, short_name=name, full_name=full_name,
                                        legal_address=address, authority=authority, phone=phone)
            session.add(organization)
            session.flush()
            organizations_created += 1
        else:
            organization.short_name = name
            organization.full_name = full_name
            organization.legal_address = address
            organization.authority = authority
            organization.phone = phone
        seen_organizations.add(organization.id)

        faculty_name = text(field(row, "Факультет"))
        if not faculty_name:
            raise ValueError(f"Строка {row_number}: не указан факультет.")
        faculty = get_or_create_faculty(session, faculty_name)
        seen_faculties.add(faculty_name)
        number = text(field(row, "Номер договора")) or "Без номера"
        contract = session.query(Contract).filter_by(organization_id=organization.id, number=number).first()
        start_date = parse_date(field(row, "Дата начала договора"))
        end_date = parse_date(field(row, "Дата окончания договора"))
        if contract is None:
            status = text(field(row, "Статус")) or "Активен"
            status = {"ACTIVE": "Активен", "CLOSED": "Закрыт"}.get(status, status)
            contract = Contract(organization_id=organization.id, number=number, status=status,
                                start_date=start_date or date.today(), end_date=end_date)
            session.add(contract)
            session.flush()
            contracts_created += 1
        else:
            if start_date:
                contract.start_date = start_date
            contract.end_date = end_date
        seen_contracts.add(contract.id)

        if not session.get(ContractFaculty, (contract.id, faculty.id)):
            session.add(ContractFaculty(contract_id=contract.id, faculty_id=faculty.id))
            session.flush()
            faculty_links_created += 1

        specialty_code = text(field(row, "Код специальности, направления специальности, специализации"))
        if specialty_code:
            qualification = text(field(row, "Квалификация"))
            specialty = get_or_create_specialty(session, specialty_code, qualification, faculty_name)
            seen_specialties.add(specialty_code)
            order = get_or_create_order(session, contract, user_id)
            item = session.query(OrderItem).filter_by(
                order_id=order.id, faculty_id=faculty.id, specialty_id=specialty.id
            ).first()
            if item is None:
                item = OrderItem(order_id=order.id, specialty_id=specialty.id,
                                 faculty_id=faculty.id, qualification_value=qualification or None)
                session.add(item)
                session.flush()
                order_items_created += 1
            else:
                item.qualification_value = qualification or None
            for year, column in years:
                try:
                    quantity = int(row[column] or 0)
                except (ValueError, TypeError) as error:
                    raise ValueError(f"Строка {row_number}, год {year}: неверная потребность.") from error
                demand = session.query(AnnualDemand).filter_by(order_item_id=item.id, year=year).first()
                if demand is None:
                    session.add(AnnualDemand(order_item_id=item.id, year=year, quantity=quantity))
                else:
                    demand.quantity = quantity
            session.flush()
        rows_processed += 1

    session.flush()
    display_name = original_filename or path.name
    result = ImportResult(rows_processed, len(seen_organizations), len(seen_contracts),
                          len(seen_faculties), len(seen_specialties), organizations_created,
                          contracts_created, faculty_links_created, order_items_created)
    current_audit_batch(session).record_values(
        AuditAction.FILE_UPLOAD, "excel_import", None, f"Импорт Excel {display_name}",
        new={
            "filename": display_name,
            "rows_processed": result.rows_processed,
            "organizations": result.organizations,
            "contracts_created": result.contracts_created,
            "order_items_created": result.order_items_created,
            "faculty_links_created": result.faculty_links_created,
        },
    )
    return result
